# -*- coding: utf-8 -*-
"""生成《毫米波雷达数据采集 SOP》Excel（执行人：周震宇 / 标注：万砚佶）。"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

OUT = r"C:/Users/tttt/WorkBuddy/大创-真正版本/docs/雷达数据采集SOP.xlsx"

# ---------- 样式 ----------
TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
SUB_FONT = Font(name="微软雅黑", size=10, color="333333")
HDR_FONT = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
CELL_FONT = Font(name="微软雅黑", size=10, color="222222")
HDR_FILL = PatternFill("solid", fgColor="2F5496")
TITLE_FILL = PatternFill("solid", fgColor="1F3864")
INFO_FILL = PatternFill("solid", fgColor="D9E1F2")
P0_FILL = PatternFill("solid", fgColor="F8CBAD")   # 红类
P1_FILL = PatternFill("solid", fgColor="FFE699")   # 黄类
P2_FILL = PatternFill("solid", fgColor="E2EFDA")   # 绿类基础
WRAP = Alignment(wrap_text=True, vertical="top", horizontal="center")
WRAP_L = Alignment(wrap_text=True, vertical="top", horizontal="left")
CTR = Alignment(vertical="center", horizontal="center")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = WRAP
        cell.border = BORDER

def style_body(ws, r0, r1, ncols, fill=None):
    for r in range(r0, r1 + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = CELL_FONT
            cell.alignment = WRAP_L if c in (3, 6, 7, 9, 15) else WRAP
            cell.border = BORDER
            if fill:
                cell.fill = fill

wb = Workbook()

# ============================================================
# Sheet 1 — 采集 SOP（周震宇），按重要性排序
# ============================================================
ws = wb.active
ws.title = "采集SOP·周震宇"
cols = ["优先级", "动作编号", "动作(中文)", "英文/枚举", "安全等级",
        "为什么重要(安全说明)", "标准做法(姿势/动作要领)", "雷达距离(m)", "朝向/站位",
        "单人段数(目标)", "每段时长(s)", "最少受试者", "总段数目标", "完成打勾(留空)", "备注"]
ncols = len(cols)

# 标题块
ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
ws.cell(1, 1, "智能康养监测系统 · 毫米波雷达数据采集 SOP（执行人：周震宇）").font = TITLE_FONT
ws.cell(1, 1).fill = TITLE_FILL
ws.cell(1, 1).alignment = CTR
ws.row_dimensions[1].height = 26

ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
ws.cell(2, 1, "雷达 IWR6843AOP ｜ 固件 xwr68xx_aop_mmw_demo.bin ｜ 采集脚本：python edge/capture_oob.py --port COMx --out data/radar/E<S>/S<subj>/A<class>/mmwave --seconds 20").font = SUB_FONT
ws.cell(2, 1).fill = INFO_FILL
ws.cell(2, 1).alignment = WRAP_L
ws.row_dimensions[2].height = 28

ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)
ws.cell(3, 1, "达标线：每类≥20段、每段≈20s、受试者≥3人(目标5人含不同年龄/身高)、总帧远超50；falling/lying_floor 必须安全模拟(垫子+慢动作+旁人保护，严禁真摔)。按“优先级”从上到下采，每完成一类在“完成打勾”列填 ✓。").font = SUB_FONT
ws.cell(3, 1).fill = INFO_FILL
ws.cell(3, 1).alignment = WRAP_L
ws.row_dimensions[3].height = 40

# 表头
for i, name in enumerate(cols, 1):
    ws.cell(5, i, name)
style_header(ws, 5, ncols)

# 数据（按重要性 P0→P3）
rows = [
    ["P0", "A08", "跌倒", "falling", "🔴红色警报",
     "项目P0核心：跌倒不报=致命。必须优先采足。",
     "站定→在厚地垫上做“失控前倾倒下”(严禁真摔！先慢动作排练2次，旁人护头)。倒地后静止3-5s再起。",
     "1.5–3.0", "正对雷达；从正前方倒下，覆盖左/中/右落点各若干段", "5–8", "20", "≥3", "≥20",
     "", "安全第一位；差动作不要凑数"],
    ["P0", "A07", "倒地不起", "lying_floor", "🔴红色警报",
     "跌倒后倒地不起=急救黄金窗口，必须识别(与卧床区分)。",
     "主动躺到地垫上保持静止(模拟跌倒后倒地不起)，侧卧/仰卧各半。",
     "1.5–3.0", "正对雷达；躺下位置覆盖左/中/右", "6–8", "20", "≥3", "≥20",
     "", "与 lying(卧床)对比采，逼模型区分"],
    ["P1", "A06", "弯腰/蹲下/拾物", "crouching", "🟡易混淆",
     "弯腰与跌倒点云极相似，混淆会致漏报→致命。最该单独成类。",
     "弯腰捡地上物品/系鞋带，保持蹲姿2-3s后起身，重复。动作慢且自然。",
     "1.5–3.0", "正对雷达；左/中/右各若干段", "8–10", "20", "≥3", "≥24",
     "", "最易和 falling 混淆，样本要多"],
    ["P2", "A05", "起身", "standing_up", "🟡易混淆",
     "起身下坠的速度特征接近跌倒，需模型区分。",
     "从椅子/床边缓慢站起→站稳，重复。",
     "1.5–3.0", "正对雷达", "8–10", "20", "≥3", "≥24",
     "", "站起瞬间有下坠速度，注意与 falling 区分"],
    ["P2", "A01", "走动", "walking", "🟢正常",
     "日常主要活动，行为识别基础。",
     "在雷达前1.5–3m范围内自然来回走(可原地踏步/踱步)。",
     "1.5–3.0", "正对+侧身各若干段", "8–10", "30", "≥3", "≥24",
     "", "覆盖不同速度(慢走/常速)"],
    ["P2", "A04", "卧床", "lying", "🟢正常",
     "卧床需与倒地不起区分，避免误报。",
     "躺在床上/沙发静止(正常休息)。",
     "1.5–3.0", "正对床", "8–10", "20", "≥3", "≥24",
     "", "人需在雷达FOV内"],
    ["P3", "A02", "站立", "standing", "🟢正常",
     "基础静态姿态。",
     "站立不动20s(可轻微晃动，模拟真实站立)。",
     "1.5–3.0", "正对雷达", "8–10", "20", "≥3", "≥24",
     "", "基础静态类"],
    ["P3", "A03", "静坐", "sitting_still", "🟢正常",
     "基础静态姿态。",
     "静坐不动20s(椅子上)。",
     "1.5–3.0", "正对雷达", "8–10", "20", "≥3", "≥24",
     "", "基础静态类"],
]
r = 6
for row in rows:
    for i, v in enumerate(row, 1):
        ws.cell(r, i, v)
    # 安全等级底色
    lvl = row[4]
    fill = P0_FILL if "红色" in lvl else (P1_FILL if "易混淆" in lvl else P2_FILL)
    style_body(ws, r, r, ncols, fill=fill)
    r += 1

# 列宽
widths = [8, 9, 13, 14, 11, 30, 40, 10, 22, 11, 10, 9, 10, 12, 16]
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A6"
# 完成打勾 下拉
dv = DataValidation(type="list", formula1='"✓,"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"N6:N{r-1}")

# ============================================================
# Sheet 2 — 标注 SOP（万砚佶）
# ============================================================
ws2 = wb.create_sheet("标注SOP·万砚佶")
ws2.merge_cells("A1:E1")
ws2.cell(1, 1, "数据标注 SOP（执行人：万砚佶）").font = TITLE_FONT
ws2.cell(1, 1).fill = TITLE_FILL
ws2.cell(1, 1).alignment = CTR
ws2.row_dimensions[1].height = 24
ws2.merge_cells("A2:E2")
ws2.cell(2, 1, "重要：标签=文件夹 A 编号，无需手写标签。重点工作是“核对一致性 + 质量分级 + 跌倒时序标注”。拿不准就发群问，别自己编。").font = SUB_FONT
ws2.cell(2, 1).fill = INFO_FILL
ws2.cell(2, 1).alignment = WRAP_L
ws2.row_dimensions[2].height = 32

ws2.cell(4, 1, "一、执行步骤").font = Font(name="微软雅黑", size=11, bold=True, color="1F3864")
steps_hdr = ["步骤", "做什么", "怎么判断/标准", "填写位置", "完成打勾(留空)"]
for i, h in enumerate(steps_hdr, 1):
    ws2.cell(5, i, h)
style_header(ws2, 5, 5)
steps = [
    ["1 收货核对", "打开 data/radar/ 目录，确认树形 E<S>/S<subj>/A<class>/mmwave/frame*.bin",
     "每类都有文件、frame 数>0", "—", ""],
    ["2 动作一致性抽查", "每类随机抽20%段，看人形点云是否对得上文件夹动作",
     "点云形态与动作相符；不符则退回周重采", "记 annotations_summary.csv 的“不合格”列", ""],
    ["3 质量分级", "每段标 好/中/差",
     "好=点云清晰人形完整；中=偶有丢帧但可用；差=人不在FOV/大量丢帧", "段目录同级 quality.txt 或汇总表", ""],
    ["4 跌倒时序标注(仅A08)", "在 A08 每段目录建 falling_events.txt，写“起始帧,结束帧”",
     "标出人开始失控下坠→触地静止的区间", "data/radar/E*/S*/A08/mmwave/falling_events.txt", ""],
    ["5 汇总", "生成 data/radar/annotations_summary.csv：动作,段数,合格段数,不合格段数,备注",
     "每类合格段数≥20 才可进训练", "annotations_summary.csv", ""],
    ["6 交付", "把 quality/事件文件随目录网盘发回，截图汇总表发群",
     "—", "—", ""],
]
r = 6
for row in steps:
    for i, v in enumerate(row, 1):
        ws2.cell(r, i, v)
    style_body(ws2, r, r, 5)
    r += 1

# 字段说明表
fr = r + 2
ws2.cell(fr, 1, "二、字段含义速查").font = Font(name="微软雅黑", size=11, bold=True, color="1F3864")
fr += 1
field_hdr = ["字段", "含义", "取值", "示例"]
for i, h in enumerate(field_hdr, 1):
    ws2.cell(fr, i, h)
style_header(ws2, fr, 4)
fields = [
    ["A编号", "动作类别，即文件夹名", "A01–A08", "A08"],
    ["S编号", "受试者编号", "S01,S02…", "S02"],
    ["E编号", "环境/场景", "E01卧室,E02客厅", "E01"],
    ["frame*.bin", "单帧点云 (N,3) float64", "雷达输出", "frame0001.bin"],
    ["quality", "段质量", "好/中/差", "中"],
    ["falling_events", "跌倒起止帧(仅A08)", "start,end", "120,180"],
    ["完成打勾", "该段标注完毕", "✓", "✓"],
]
rr = fr + 1
for row in fields:
    for i, v in enumerate(row, 1):
        ws2.cell(rr, i, v)
    style_body(ws2, rr, rr, 4)
    rr += 1

for i, w in enumerate([20, 42, 30, 34], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A6"

# ============================================================
# Sheet 3 — 空白采集记录表（带下拉）
# ============================================================
ws3 = wb.create_sheet("采集记录表(空白)")
rec_hdr = ["日期", "受试者(Sxx)", "动作(Axx)", "段号", "距离(m)", "帧数", "质量(好/中/差)", "标注完成(留空)", "备注"]
for i, h in enumerate(rec_hdr, 1):
    ws3.cell(1, i, h)
style_header(ws3, 1, len(rec_hdr))
for r in range(2, 32):
    for c in range(1, len(rec_hdr) + 1):
        cell = ws3.cell(r, c)
        cell.font = CELL_FONT
        cell.border = BORDER
        cell.alignment = WRAP
dvq = DataValidation(type="list", formula1='"好,中,差"', allow_blank=True)
ws3.add_data_validation(dvq)
dvq.add("G2:G31")
dvd = DataValidation(type="list", formula1='"✓,"', allow_blank=True)
ws3.add_data_validation(dvd)
dvd.add("H2:H31")
for i, w in enumerate([12, 12, 11, 7, 9, 8, 14, 14, 20], 1):
    ws3.column_dimensions[get_column_letter(i)].width = w
ws3.freeze_panes = "A2"

wb.save(OUT)
print("SAVED:", OUT)
